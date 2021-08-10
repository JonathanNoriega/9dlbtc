from tkinter import *
from requests import *
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from time import *
from infi.systray import SysTrayIcon
from binance.client import Client
from py_currency_converter import convert
import json
import time
import ctypes
import os
import sys

api_key = os.environ.get('binance_api')
api_secret = os.environ.get('binance_secret')

#btc price

amount = 0.0001273

client = Client(api_key, api_secret)
btc_price = client.get_symbol_ticker(symbol="BTCUSDT")
usd_price = float(btc_price["price"])

price = convert(base='USD', amount=usd_price, to=['COP'])
print(price['COP'])     

#Excel

book = load_workbook("log.xlsx")
sheet = book.active
 
def bye(sysTrayIcon):
    os.system("taskkill /f /im  python.exe")

menu_options = (())
systray = SysTrayIcon("../9dbtc/img/btc.ico", "9dbtc", menu_options, on_quit=bye)
systray.start()

def save_data(delay):
    saving = True
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 25
    sheet.column_dimensions["C"].width = 25
    while True:
        client = Client(api_key, api_secret)
        btc_price = client.get_symbol_ticker(symbol="BTCUSDT")
        usd_price = float(btc_price["price"])
        price = convert(base='USD', amount=usd_price, to=['COP'])
        d9 = amount * float(price['COP']) / 1
        
        sheet['A1'] = "Date"
        sheet['B1'] = "Price"
        sheet['C1'] = "Your bitcoin is worth"
        lastRowA = 'A' + str(sheet.max_row + 1)
        lastRowB = 'B' + str(sheet.max_row + 1)
        lastRowC = 'C' + str(sheet.max_row + 1)
        sheet[lastRowA] = strftime("%a, %d %b %Y %H:%M:%S")
        sheet[lastRowB] = price['COP']
        sheet[lastRowC] = d9
        book.save('log.xlsx')
        time.sleep(delay)

print("#####################################")
print("#                                   #")
print("#                                   #")
print("#     Saving data in 'log.xlsx'     #")
print("#                                   #")
print("#                                   #")
print("#####################################")

save_data(1800)
